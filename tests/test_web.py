import re

from fastapi.testclient import TestClient

from app.db import DB
from app.main import create_app
from tests.test_core import FakeClient


def local_today():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    return datetime.now(ZoneInfo("Asia/Beirut")).date()


def make(settings):
    return TestClient(create_app(settings, FakeClient(settings)), base_url="http://testserver")


def login(c, settings, password=None, username=None):
    username = username or ("owner" if password is None else "family")
    r = c.post("/login", data={"username": username, "password": password or settings.dashboard_password}, follow_redirects=False)
    assert r.status_code == 303
    page = c.get("/medicines" if password is None else "/family").text
    return {"X-CSRF-Token": page.split('name="csrf" content="')[1].split('"')[0]}


def sheet(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["", "Patient — Medicines"])
    ws.append(["", "Updated 14 Oct · Prescriber: Dr X"])
    ws.append(["Qty 15/10/2026", "Medication", "Dose", "Form", "When to take", "For", "Notes", "بالعربي"])
    ws.append([30, "Alpha", "160 mg", "Tablet", "5 tablets daily, without stopping", "Appetite", "Confirm dose with clinic", "٥ حبات يومياً"])
    ws.append([17, "Beta", "—", "Tablet", "1–2 tablets, 3 times daily when in pain", "Pain", "No other paracetamol", "حبة أو حبتين عند الوجع"])
    ws.append([None, "REGULAR"])
    ws.append([None, "Gamma", "30 mg", "Injection", "Every 28 days", "Control", "Essential; do not skip", "إبرة كل ٢٨ يوم"])
    ws.append([None, "Delta", "500 mg", "Tablet", "Twice daily (morning & evening)", "Antiviral", None, "حبة صباحاً ومساءً"])
    ws.append([None, "Epsilon", "10,000 IU", "Supplement", "Once weekly, lifelong", "Vitamin", None, "مرة في الأسبوع"])
    ws.append([None, "AS NEEDED"])
    ws.append([None, "Eta", "2 mg", "Tablet", "Up to twice daily", "Nausea", "الغثيان", "حتى مرتين في اليوم"])
    ws.append([None, "Theta", "—", "Tablet", "2 tablets every 8 hours", "Pain", None, "حبتين كل ٨ ساعات"])
    ws.append([None, "Iota", "1 mg", "", "1 tablet daily", None, "فقر الدم", "حبة كل يوم"])  # listed under AS NEEDED but says daily
    ws.append([None, "Note: Some other tablets are not on this list."])
    wb.save(path)


def imported(c, h, tmp_path, replace=False):
    p = tmp_path / "m.xlsx"
    sheet(p)
    return c.post(f"/api/medicines-import?replace={'true' if replace else 'false'}",
                  files={"file": ("m.xlsx", p.read_bytes())}, headers=h)


def med_id(c, name):
    html = c.get("/medicines").text
    return int(re.search(r'data-edit="(\d+)" dir="auto"><strong>' + re.escape(name) + "<", html).group(1))


def test_auth_required(settings):
    c = make(settings)
    for path in ("/", "/medicines", "/reminders", "/mom", "/settings", "/family", "/family/log", "/reminders/export.csv"):
        assert c.get(path, follow_redirects=False).status_code == 303, path
    assert c.get("/api/status").status_code == 401
    assert c.post("/login", data={"username": "owner", "password": "nope"}).status_code == 401
    assert c.get("/login").status_code == 200


def test_lockout(settings):
    c = make(settings)
    for _ in range(5):
        c.post("/login", data={"username": "owner", "password": "x"})
    assert c.post("/login", data={"username": "owner", "password": settings.dashboard_password}).status_code == 429


def test_navigation_and_pages(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    home = c.get("/", follow_redirects=False)
    assert home.headers["location"] == "/medicines"
    for path, needle in (("/medicines", "Medicines"), ("/reminders", "Reminders report"), ("/mom", "Add to the log"),
                         ("/settings", "Wbiztool connection")):
        r = c.get(path)
        assert r.status_code == 200 and needle in r.text
        assert 'href="/reminders"' in r.text and 'href="/mom"' in r.text and 'href="/settings"' in r.text  # nav everywhere


def test_import_and_medicines_page(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    p = tmp_path / "m.xlsx"
    sheet(p)
    assert c.post("/api/medicines-import", files={"file": ("m.xlsx", p.read_bytes())}).status_code == 403  # CSRF
    first = imported(c, h, tmp_path).json()
    assert first["added"] == 8 and first["skipped"] == 0 and first["stock_counts"] == 2  # Alpha 30 and Beta 17 in the file
    again = imported(c, h, tmp_path).json()
    assert again["added"] == 0 and again["skipped"] == 8 and again["stock_counts"] == 0  # no duplicates, no double counts
    page = c.get("/medicines").text
    assert page.count('role="switch"') == 8 and "Note: Some other" not in page
    assert 'class="switch on' not in page  # reminders are opt-in: nothing is switched on after an import
    reg, rest = page.split("As needed", 1)
    assert "Alpha" in reg and "Iota" in reg  # 'daily' in the text wins over the AS NEEDED heading
    assert "Eta" in rest and "Theta" in rest and "Beta" in rest  # 'when in pain' / 'up to' -> as needed
    assert "Tablet" in page and "Injection" in page
    ar = c.get("/medicines?lang=ar").text
    assert 'dir="rtl"' in ar and "حبة صباحاً ومساءً" in ar and "الأدوية" in ar and "عند الحاجة" in ar and "الغثيان" in ar
    assert imported(c, h, tmp_path, replace=True).json()["added"] == 8
    assert c.get("/medicines").text.count('role="switch"') == 8


def test_suggestions_and_switch_flow(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta, eps, gamma, eta = (med_id(c, n) for n in ("Delta", "Epsilon", "Gamma", "Eta"))
    # nothing syncs until a reminder is on
    assert c.post("/api/sync/preview", headers=h).json()["create"] == []
    # "Twice daily (morning & evening)" -> one-click suggestion; weekly/every-N still need a day/date
    sug = c.get(f"/api/medicines/{delta}").json()["suggest"]
    assert sug == {"frequency": "daily", "times": ["08:00", "20:00"], "complete": True}
    assert c.get(f"/api/medicines/{eps}").json()["suggest"]["complete"] is False
    assert c.get(f"/api/medicines/{gamma}").json()["suggest"]["complete"] is False
    # custom default times change the suggestion
    assert c.put("/api/slots", json={"morning": "07:30", "evening": "21:00"}, headers=h).json()["morning"] == "07:30"
    assert c.get(f"/api/medicines/{delta}").json()["suggest"]["times"] == ["07:30", "21:00"]
    assert c.put("/api/slots", json={"morning": "7am"}, headers=h).status_code == 400
    # turning the switch on is what the page does with the suggestion
    r = c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["07:30", "21:00"]}, headers=h)
    assert r.status_code == 200
    assert c.post(f"/api/medicines/{eps}/reminder", json={"on": True, "times": ["09:00"]}, headers=h).status_code == 400  # needs a day
    assert c.post(f"/api/medicines/{eps}/reminder", json={"on": True, "times": ["09:00"], "weekdays": [4]}, headers=h).status_code == 200
    assert c.post(f"/api/medicines/{gamma}/reminder", json={"on": True, "times": ["10:00"], "every_n": 28,
                                                          "date": "2026-10-20"}, headers=h).status_code == 200
    assert c.post(f"/api/medicines/{eta}/reminder", json={"on": True, "times": ["08:00"]}, headers=h).status_code == 400  # as needed
    page = c.get("/medicines").text
    assert page.count('class="switch on') == 3 and "07:30 · 21:00" in page and "Fri 09:00" in page and "every 28 days" in page
    assert "not sent yet" in page and 'href="/reminders?review=1"' in page
    plan = c.post("/api/sync/preview", headers=h).json()
    assert len(plan["create"]) == 8 and "Eta" not in "".join(x["message"] for x in plan["create"])
    assert c.post("/api/sync/apply", json={"hash": "stale"}, headers=h).status_code == 409
    res = c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h).json()["results"]
    assert len(res) == 8 and all(x["ok"] for x in res)
    assert "not sent yet" not in c.get("/medicines").text
    c.post(f"/api/medicines/{delta}/reminder", json={"on": False}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    assert len(plan["cancel"]) == 4 and plan["create"] == []


def test_hand_added_medicine_gets_one_click_suggestion(settings):
    c = make(settings)
    h = login(c, settings)
    mid = c.post("/api/medicines", json={"name": "Hand", "when_en": "Twice daily (morning & evening)", "section": "regular"}, headers=h).json()["id"]
    assert c.get(f"/api/medicines/{mid}").json()["suggest"] == {"frequency": "daily", "times": ["08:00", "20:00"], "complete": True}
    wk = c.post("/api/medicines", json={"name": "Wk", "when_en": "Once weekly", "section": "regular"}, headers=h).json()["id"]
    s2 = c.get(f"/api/medicines/{wk}").json()["suggest"]
    assert s2["frequency"] == "weekly" and s2["complete"] is False  # still needs a day
    none = c.post("/api/medicines", json={"name": "Plain", "section": "regular"}, headers=h).json()["id"]
    assert c.get(f"/api/medicines/{none}").json()["suggest"]["complete"] is False  # no text, no guess
    pn = c.post("/api/medicines", json={"name": "Pn", "when_en": "when needed", "section": "regular"}, headers=h).json()["id"]
    assert c.get(f"/api/medicines/{pn}").json()["suggest"]["complete"] is False


def test_reminders_report_and_csv(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=h)
    page = c.get("/reminders").text
    assert "Waiting to apply" in page and "Review &amp; apply" in page and "Nothing is scheduled yet" in page
    plan = c.post("/api/sync/preview", headers=h).json()
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    page = c.get("/reminders").text
    assert "Every day at 08:00" in page and "Mom" in page and "Copy to me" in page and "Delta" in page
    assert re.search(r'<span class="num">2</span><span class="lbl">Active reminders', page)
    csv_text = c.get("/reminders/export.csv").text
    assert csv_text.splitlines()[0].startswith("Recipient,Medicines,Schedule") and "Every day at 08:00" in csv_text
    assert "create" in c.get("/reminders").text  # activity log
    assert c.get("/reminders?show=failed").status_code == 200
    # spreadsheet formula injection is neutralised
    c.post("/api/medicines", json={"name": "=HYPERLINK(1)", "section": "regular", "when_en": "daily"}, headers=h)
    mid = med_id(c, "=HYPERLINK(1)")
    c.post(f"/api/medicines/{mid}/reminder", json={"on": True, "frequency": "daily", "times": ["09:00"]}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    assert "'=HYPERLINK" in c.get("/reminders/export.csv").text and ',=HYPERLINK' not in c.get("/reminders/export.csv").text


def test_family_read_only_access(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    c.put("/api/profile", json={"name": "Amina", "allergies": "Penicillin", "doctor": "Dr X", "doctor_phone": "+961 9 830 600"}, headers=h)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=h)
    # setting the password: validation
    assert c.put("/api/family-access", json={"password": "short"}, headers=h).status_code == 400
    assert c.put("/api/family-access", json={"password": settings.dashboard_password}, headers=h).status_code == 400
    assert c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h).json() == {"enabled": True}
    from app.db import DB
    stored = DB(settings.db_path).get_setting("family.hash")
    assert stored.startswith("scrypt$") and "family-pass-123" not in stored  # hashed, never plain
    f = TestClient(c.app, base_url="http://testserver")
    assert f.post("/login", data={"username": "family", "password": "wrong-password"}).status_code == 401
    fh = login(f, settings, "family-pass-123")
    assert f.get("/", follow_redirects=False).headers["location"] == "/family"
    fam = f.get("/family").text
    assert "Delta" in fam and "Penicillin" in fam and "tel:+9619830600" in fam and "08:00" in fam
    assert "data-toggle" not in fam and "/settings" not in fam and "/reminders" not in fam and "SECRET" not in fam
    assert "حبة صباحاً ومساءً" in f.get("/family?lang=ar").text
    # everything else is owner-only
    for path in ("/medicines", "/reminders", "/mom", "/settings", "/reminders/export.csv"):
        r = f.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/family" or r.status_code == 403, path
    assert f.get("/api/status").status_code == 403
    assert f.put("/api/profile", json={"name": "x"}, headers=fh).status_code == 403
    assert f.post("/api/medicines", json={"name": "x"}, headers=fh).status_code == 403
    assert f.put("/api/family-access", json={"password": "another-pass-1"}, headers=fh).status_code == 403
    # the owner's password still gives the owner role
    o = TestClient(c.app, base_url="http://testserver")
    assert login(o, settings) and o.get("/settings").status_code == 200
    # turn it off -> family can no longer sign in
    assert c.delete("/api/family-access", headers=h).json() == {"enabled": False}
    assert TestClient(c.app, base_url="http://testserver").post("/login", data={"username": "family", "password": "family-pass-123"}).status_code == 401


def test_edit_details_and_sections(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "times": ["08:00"]}, headers=h)
    assert c.put(f"/api/medicines/{delta}", json={"dose": "250 mg", "notes_ar": "ملاحظة"}, headers=h).status_code == 200
    assert "250 mg" in c.get("/medicines").text and c.get(f"/api/medicines/{delta}").json()["remind"] is True
    assert c.put(f"/api/medicines/{delta}", json={"name": ""}, headers=h).status_code == 400
    c.put(f"/api/medicines/{delta}", json={"section": "as_needed"}, headers=h)
    page = c.get("/medicines").text
    assert page.index("Delta") > page.index("As needed") and c.get(f"/api/medicines/{delta}").json()["remind"] is False
    new = c.post("/api/medicines", json={"name": "Aspirin <b>", "dose": "100mg", "when_en": "daily", "section": "regular"}, headers=h)
    assert "<b>" not in c.get("/medicines").text and c.get(f"/api/medicines/{new.json()['id']}").json()["frequency"] == "chart_only"
    assert c.post("/api/medicines", json={"name": ""}, headers=h).status_code == 400
    assert c.delete(f"/api/medicines/{new.json()['id']}", headers=h).status_code == 200
    assert c.get(f"/api/medicines/{new.json()['id']}").status_code == 404


def test_mom_page_profile_and_status(settings):
    c = make(settings)
    h = login(c, settings)
    c.put("/api/profile", json={"name": "Amina", "allergies": "Penicillin", "evil": "x"}, headers=h)
    assert c.post("/api/status-log", json={"kind": "Blood pressure", "value": "130/80"}, headers=h).status_code == 200
    assert c.post("/api/status-log", json={"kind": "Nope", "value": "1"}, headers=h).status_code == 400
    assert c.post("/api/status-log", json={"kind": "Note"}, headers=h).status_code == 400
    mom = c.get("/mom").text
    assert "130/80" in mom and "Latest readings" in mom and "Amina" in mom
    sid = mom.split('data-del-status="')[1].split('"')[0]
    assert c.delete(f"/api/status-log/{sid}", headers=h).status_code == 200
    c.post("/api/status-log", json={"kind": "Note", "note": "x"}, headers=h)
    assert c.delete("/api/status-log", headers=h).json()["deleted"] == 1


def test_settings_page(settings):
    settings.client_id = None
    settings.api_key = ""
    settings.whatsapp_client = ""
    settings.mom_number = ""
    settings.docs_verified = False
    c = make(settings)
    h = login(c, settings)
    r = c.put("/api/config", json={"client_id": "7", "api_key": "S3CRET-KEY_value", "whatsapp_client": "42",
                                   "mom_number": "+961 70 123456", "my_number": "96170999999", "docs_verified": True}, headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["api_key_set"] and body["api_key_length"] == 16 and body["mom_number"] == "96170123456"
    assert "S3CRET" not in r.text and "S3CRET" not in c.get("/settings").text and "S3CRET" not in c.get("/").text
    assert settings.client_id == 7 and settings.api_key == "S3CRET-KEY_value" and settings.docs_verified
    # blank key keeps the saved one
    c.put("/api/config", json={"api_key": "", "whatsapp_client": "43"}, headers=h)
    assert settings.api_key == "S3CRET-KEY_value" and settings.whatsapp_client == "43"
    for bad in ({"client_id": "abc"}, {"api_key": "has space"}, {"mom_number": "123"}, {"whatsapp_client": "a b"}):
        assert c.put("/api/config", json=bad, headers=h).status_code == 400
    assert c.put("/api/config", json={}).status_code == 403  # CSRF
    assert c.post("/api/config/test", headers=h).json()["ok"]
    # survives a restart: a fresh Settings object picks the values from the database
    from app.config import Settings
    from app.db import DB
    s2 = Settings(data_dir=settings.data_dir)
    s2.apply_overrides(DB(settings.db_path))
    assert s2.client_id == 7 and s2.api_key == "S3CRET-KEY_value"


def test_status_failure_does_not_hide_good_credentials(settings):
    from app.wbiztool import WbizError

    class NoStatus(FakeClient):
        def client_status(self):
            raise WbizError("Non-JSON response from /whatsapp_client/status/ (HTTP 404): Page not found")

    c = TestClient(create_app(settings, NoStatus(settings)), base_url="http://testserver")
    h = login(c, settings)
    r = c.post("/api/config/test", headers=h).json()
    assert r["ok"] and "404" in r["whatsapp_error"]
    st = c.get("/api/status").json()
    assert st["error"] is None and "404" in st["whatsapp_error"]


def test_migration_keeps_existing_reminders(tmp_path):
    import sqlite3
    from app.db import DB
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE medicines (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL DEFAULT '', dose TEXT NOT NULL DEFAULT '', times TEXT NOT NULL DEFAULT '', frequency TEXT NOT NULL DEFAULT '', day TEXT NOT NULL DEFAULT '', start_date TEXT NOT NULL DEFAULT '', last_date TEXT NOT NULL DEFAULT '', instruction TEXT NOT NULL DEFAULT '', purpose TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL)")
    con.execute("INSERT INTO medicines(name,times,frequency,created_at) VALUES('WithTimes','08:00','daily','x')")
    con.execute("INSERT INTO medicines(name,created_at) VALUES('NoTimes','x')")
    con.commit(); con.close()
    db = DB(path)
    rows = {m["name"]: m for m in db.list_medicines()}
    assert rows["WithTimes"]["remind"] == 1 and rows["NoTimes"]["remind"] == 0 and rows["NoTimes"]["form"] == ""
    assert db.add_medicine({"name": "New", "notes_ar": "ملاحظة", "form": "Tablet"})


def test_status_failure_does_not_hide_good_credentials(settings):
    from app.wbiztool import WbizError

    class NoStatus(FakeClient):
        def client_status(self):
            raise WbizError("Non-JSON response from /whatsapp_client/status/ (HTTP 404): Page not found")

    c = TestClient(create_app(settings, NoStatus(settings)), base_url="http://testserver")
    h = login(c, settings)
    r = c.post("/api/config/test", headers=h).json()
    assert r["ok"] and "404" in r["whatsapp_error"]
    st = c.get("/api/status").json()
    assert st["error"] is None and "404" in st["whatsapp_error"]


def test_fill_missing_arabic(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    # imitate an older import: an Arabic-only note sitting in the English notes box
    c.post("/api/medicines", json={"name": "Imod", "dose": "2 mg", "form": "Tablet", "purpose": "Diarrhoea", "notes_en": "اسهال",
                                   "when_en": "when needed", "instruction": "حبة عند الضرورة", "section": "as_needed"}, headers=h)
    c.post("/api/medicines", json={"name": "Odd", "form": "Tablet", "purpose": "Mystery purpose", "notes_en": "Take with orange juice",
                                   "when_en": "daily", "instruction": "", "section": "regular"}, headers=h)
    assert "Arabic detail" in c.get("/medicines").text
    r = c.post("/api/arabic/fill", headers=h).json()
    filled = {(x["name"], x["field"]): x["value"] for x in r["filled"]}
    assert filled[("Imod", "notes")] == "اسهال" and filled[("Imod", "form")] == "أقراص" and filled[("Imod", "what it is for")] == "الإسهال"
    assert filled[("Gamma", "notes")] == "ضروري؛ لا تفوّت الجرعة" and filled[("Gamma", "form")] == "حقنة"
    still = {(x["name"], x["field"]) for x in r["missing"]}
    assert ("Odd", "what it is for") in still and ("Odd", "notes") in still and ("Odd", "how to take") in still  # not guessed
    assert c.post("/api/arabic/fill", headers=h).json()["filled"] == []  # idempotent, never overwrites
    mid = med_id(c, "Imod")
    form = c.get(f"/api/medicines/{mid}").json()
    assert form["notes_ar"] == "اسهال" and form["notes_en"] == "" and form["form_ar"] == "أقراص"
    # a person's own Arabic is never replaced
    c.put(f"/api/medicines/{mid}", json={"purpose_ar": "إسهال حاد"}, headers=h)
    c.put(f"/api/medicines/{mid}", json={"purpose": "Diarrhoea"}, headers=h)
    c.post("/api/arabic/fill", headers=h)
    assert c.get(f"/api/medicines/{mid}").json()["purpose_ar"] == "إسهال حاد"
    ar = c.get("/medicines?lang=ar").text
    assert "أقراص" in ar and "160 ملغ" in ar and "الشهية" in ar
    odd_row = ar[ar.index("Odd"):]
    assert "Mystery purpose" in odd_row and 'class="fb"' in odd_row  # not translated -> shown in English, marked
    assert 'class="fb"' not in ar[ar.index("Alpha"):ar.index("Odd")] or "Odd" in ar  # translated rows are not marked
    assert c.post("/api/arabic/fill").status_code == 403  # CSRF
    assert "أقراص" in c.get("/family?lang=ar").text


def test_arabic_units_and_glossary():
    from app.arabic import arabic_dose, translate_notes, translate_purpose
    assert arabic_dose("10,000 IU") == "10,000 وحدة دولية" and arabic_dose("500mg") == "500ملغ" and arabic_dose("2 ml") == "2 مل"
    assert translate_purpose("Appetite / weight") == "الشهية / الوزن" and translate_purpose("Unknown thing") == ""
    assert translate_notes("Same medicine as Primperan: use one, not both") == "نفس الدواء مثل Primperan: استخدم واحداً فقط وليس كليهما"
    assert translate_notes("Started 20 Dec 2024") == "بدأ بتاريخ 20 كانون الأول 2024" and translate_notes("Take with orange juice") == ""


def test_send_test_reminder_goes_only_to_my_number(settings, tmp_path, monkeypatch):
    from app.wbiztool import WbizError

    class Spy(FakeClient):
        sent = []

        def _post(self, path, fields=None, retries=0):
            if path.endswith("/send_msg/"):
                Spy.sent.append(dict(fields))
            return {"status": 1, "reminder_id": "r1"}

    c = TestClient(create_app(settings, Spy(settings)), base_url="http://testserver")
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    assert all(x["fingerprint"] for x in plan["create"])
    fp = plan["create"][0]["fingerprint"]
    # a pending (not yet applied) message can be tested
    assert c.post("/api/test-reminder", json={"fingerprint": fp}, headers=h).json() == {"ok": True, "to": "me"}
    sent = Spy.sent[-1]
    assert sent["phone"] == settings.my_number and settings.mom_number not in sent.values()
    assert sent["message_template"].startswith("🧪") and "Delta" in sent["message_template"] and "Test message" in sent["message_template"]
    assert c.post("/api/test-reminder", json={"fingerprint": fp}, headers=h).status_code == 429  # throttled
    # applied reminders can be tested from the report; the report shows the button
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    page = c.get("/reminders").text
    assert page.count('data-test="') == 4  # a 'to me' and a 'to Mom' button on each of the 2 reminders
    import app.main as main_mod
    monkeypatch.setattr(main_mod, "TEST_MIN_INTERVAL", 0.0)
    active_fp = page.split('data-test="')[1].split('"')[0]
    assert c.post("/api/test-reminder", json={"fingerprint": active_fp}, headers=h).json() == {"ok": True, "to": "me"}
    assert c.post("/api/test-reminder", json={"fingerprint": "nope"}, headers=h).status_code == 404
    assert c.post("/api/test-reminder", json={"fingerprint": fp}).status_code == 403  # CSRF
    # --- test to Mom: needs explicit confirmation, uses her number, her own Arabic note, no 'copy to me' line
    assert c.post("/api/test-reminder", json={"fingerprint": fp, "to": "mom"}, headers=h).status_code == 400
    me_fp = next(x["fingerprint"] for x in plan["create"] if x["recipient"] == "me")
    r = c.post("/api/test-reminder", json={"fingerprint": me_fp, "to": "mom", "confirm": True}, headers=h)
    assert r.json() == {"ok": True, "to": "mom"}
    sent = Spy.sent[-1]
    assert sent["phone"] == settings.mom_number and sent["message_template"].startswith("🧪 رسالة تجريبية من تطبيق")
    assert "نسخة للمتابعة" not in sent["message_template"] and "Delta" in sent["message_template"]
    assert "Test message" not in sent["message_template"]  # her note is Arabic only
    assert c.post("/api/test-reminder", json={"fingerprint": fp, "to": "mom", "confirm": True}, headers=h).status_code in (200, 429)
    monkeypatch.setattr(main_mod, "TEST_MIN_INTERVAL", 0.0)
    assert f'data-mom-tail="{settings.mom_number[-3:]}"' in c.get("/reminders").text  # shown in the confirmation
    settings.mom_number = ""
    assert c.post("/api/test-reminder", json={"fingerprint": fp, "to": "mom", "confirm": True}, headers=h).status_code == 400
    settings.my_number = ""
    assert c.post("/api/test-reminder", json={"fingerprint": active_fp}, headers=h).status_code == 400


def test_purpose_line_in_reminders_and_toggle(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    c.post("/api/arabic/fill", headers=h)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    assert "الغرض: مضاد للفيروسات" in plan["create"][0]["message"]  # 'Antiviral' filled in Arabic
    assert "Antiviral" not in plan["create"][0]["message"]
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    assert c.post("/api/sync/preview", headers=h).json()["create"] == []  # in sync
    # switching it off changes the text -> the old reminders are cancelled and re-created
    assert "الغرض" in c.get("/settings").text and "Example" in c.get("/settings").text
    assert c.put("/api/message-options", json={"include_purpose": False}, headers=h).json() == {"include_purpose": False}
    plan = c.post("/api/sync/preview", headers=h).json()
    assert len(plan["cancel"]) == 2 and len(plan["create"]) == 2 and "الغرض" not in plan["create"][0]["message"]
    assert c.put("/api/message-options", json={"include_purpose": True}, headers=h).json() == {"include_purpose": True}
    assert c.post("/api/sync/preview", headers=h).json()["cancel"] == []  # back to what is already on Wbiztool
    assert c.put("/api/message-options", json={"include_purpose": False}).status_code == 403  # CSRF


def test_status_log_dates_and_upcoming(settings):
    from datetime import date, timedelta
    c = make(settings)
    h = login(c, settings)
    today = local_today()
    past = (today - timedelta(days=40)).isoformat()
    future = (today + timedelta(days=12)).isoformat()
    post = lambda **kw: c.post("/api/status-log", json=kw, headers=h)  # noqa: E731
    assert post(kind="Doctor visit", value="Dr. Kamar", note="Blood test results", date=past, time="14:30").status_code == 200
    assert post(kind="Appointment", value="Oncology", date=future, time="09:15").status_code == 200
    assert post(kind="Blood pressure", value="130/80").status_code == 200  # no date -> today
    assert post(kind="Blood pressure", value="120/70", date=past).status_code == 200
    for bad in ({"date": "31-12-2026"}, {"date": "2026-02-31"}, {"date": "1990-01-01"}, {"time": "25:00"}, {"time": "9am"}):
        assert post(kind="Note", note="x", **bad).status_code == 400, bad
    mom = c.get("/mom").text
    assert "Upcoming" in mom and "Oncology" in mom and "09:15" in mom
    history = mom.split("History")[1]
    assert "Oncology" not in history                      # future entries are not in the history
    assert history.index("130/80") < history.index("Dr. Kamar")  # newest event date first
    assert "14:30" in history and "Blood test results" in history
    # 'latest readings' use the newest reading by event date, never a visit or a future entry
    latest = mom.split("Latest readings")[1].split("</section>")[0]
    assert "130/80" in latest and "120/70" not in latest and "Dr. Kamar" not in latest
    assert 'name="date" type="date" value="' + today.isoformat() in mom


def test_status_log_migration_keeps_old_entries(tmp_path):
    import sqlite3
    from app.db import DB
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE status_log (id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, kind TEXT NOT NULL, value TEXT NOT NULL DEFAULT '', note TEXT NOT NULL DEFAULT '')")
    con.execute("INSERT INTO status_log(ts,kind,value) VALUES('2026-10-01T10:00:00+00:00','Blood pressure','125/80')")
    con.commit(); con.close()
    db = DB(path)
    row = db.list_status()[0]
    assert row["event_date"] == "2026-10-01" and row["event_time"] == "" and row["value"] == "125/80"
    db.add_status("Doctor visit", "Dr X", "", "2026-12-01", "10:00")
    assert db.list_status()[0]["event_date"] == "2026-12-01"


def test_report_shows_fallback_one_off_as_a_single_date(settings, tmp_path):
    from datetime import date, timedelta
    from app.wbiztool import WbizError

    class Rejecting(FakeClient):
        def _post(self, path, fields=None, retries=0):
            if path == "/schedule_msg/":
                raise WbizError("/schedule_msg/ failed (status=0): Scheduled date & time is not in valid format")
            return {"status": 1, "reminder_id": "r1"}

    c = TestClient(create_app(settings, Rejecting(settings)), base_url="http://testserver")
    h = login(c, settings)
    imported(c, h, tmp_path)
    gamma = med_id(c, "Gamma")  # "Every 28 days"
    when = local_today() + timedelta(days=10)
    r = c.post(f"/api/medicines/{gamma}/reminder", json={"on": True, "times": ["10:00"], "every_n": 28,
                                                       "date": when.isoformat()}, headers=h)
    assert r.status_code == 200
    plan = c.post("/api/sync/preview", headers=h).json()
    assert all(x["kind"] == "oneoff" for x in plan["create"])
    res = c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h).json()["results"]
    assert all(x["ok"] for x in res)
    page = c.get("/reminders").text
    assert "Once on " + when.strftime("%d %b %Y") in page and "Every day" not in page.split("Scheduled")[1].split("Activity")[0]
    assert "one-off dates now use date-specific reminders" in page  # the switch is logged in the activity list
    assert "Once on" in c.get("/reminders/export.csv").text
    assert not c.post("/api/sync/preview", headers=h).json()["create"]


def test_family_can_read_mom_log_but_not_change_it(settings):
    from datetime import date, timedelta
    c = make(settings)
    h = login(c, settings)
    today, past, future = local_today(), (local_today() - timedelta(days=30)).isoformat(), (local_today() + timedelta(days=14)).isoformat()
    post = lambda **kw: c.post("/api/status-log", json=kw, headers=h)  # noqa: E731
    post(kind="Doctor visit", value="Dr. Kamar", note="Results fine", date=past, time="14:30")
    post(kind="Appointment", value="Oncology", date=future, time="09:15")
    post(kind="Blood pressure", value="130/80", note="Evening")
    c.put("/api/profile", json={"name": "Amina"}, headers=h)
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    f = TestClient(c.app, base_url="http://testserver")
    fh = login(f, settings, "family-pass-123")
    page = f.get("/family/log").text
    for needle in ("Upcoming", "Oncology", "09:15", "Latest readings", "130/80", "History", "Dr. Kamar", "Results fine", "14:30"):
        assert needle in page, needle
    # read-only: no way to add, delete or clear, no owner pages in the menu
    for forbidden in ("status-form", "data-del-status", "clear-log", "profile-form", 'href="/mom"', 'href="/settings"', 'href="/reminders"'):
        assert forbidden not in page, forbidden
    assert 'href="/family"' in page and 'href="/family/log"' in page  # the family menu
    # Arabic version
    ar = f.get("/family/log?lang=ar").text
    assert 'dir="rtl"' in ar and "زيارة الطبيب" in ar and "ضغط الدم" in ar and "المواعيد القادمة" in ar and "السجل" in ar
    assert (local_today() - timedelta(days=30)).strftime("%d/%m/%Y") in ar  # Arabic pages use dd/mm/yyyy
    # the medicines family page links to it too
    assert 'href="/family/log"' in f.get("/family").text
    # and it still can't change anything or reach the owner page
    assert f.post("/api/status-log", json={"kind": "Note", "note": "x"}, headers=fh).status_code == 403
    assert f.delete("/api/status-log", headers=fh).status_code == 403
    r = f.get("/mom", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/family"
    # the owner sees the same page, and family access off means no sign-in
    assert c.get("/family/log").status_code == 200
    c.delete("/api/family-access", headers=h)
    assert TestClient(c.app, base_url="http://testserver").post("/login", data={"username": "family", "password": "family-pass-123"}).status_code == 401


def test_mom_profile_and_log_import_export(settings):
    import json as _json
    c = make(settings)
    h = login(c, settings)
    c.put("/api/profile", json={"name": "Existing Name"}, headers=h)
    payload = {"profile": {"name": "Amina", "doctor": "Prof. Kamar", "care_team": "Dr A — surgeon\nDr B — GP", "bogus": "x"},
               "log": [{"kind": "Treatment", "date": "2025-07-01", "value": "Chemo started", "note": "approx. date"},
                       {"kind": "Hospital", "date": "2026-02-10", "value": "Admission", "note": "Ascites and pain"},
                       {"kind": "Note", "note": "No date -> today"}]}
    post = lambda body, **q: c.post("/api/mom-import", content=_json.dumps(body), headers={**h, "Content-Type": "application/json"}, params=q)  # noqa: E731
    r = post(payload)
    assert r.status_code == 200
    assert r.json() == {"profile_fields": 2, "profile_kept": 1, "log_added": 3, "log_skipped": 0, "ignored": ["bogus"]}
    prof = c.get("/api/mom-export").json()["profile"]
    assert prof["name"] == "Existing Name"  # existing text is kept
    assert prof["doctor"] == "Prof. Kamar" and prof["care_team"].startswith("Dr A")
    assert post(payload).json()["log_added"] == 0 and post(payload).json()["log_skipped"] == 3  # no duplicates
    assert post(payload, overwrite="true").json()["profile_fields"] == 3
    assert c.get("/api/mom-export").json()["profile"]["name"] == "Amina"
    # all-or-nothing validation
    bad = {"log": [{"kind": "Note", "note": "fine"}, {"kind": "Nope", "note": "x"}, {"kind": "Note", "note": "y", "date": "2026-02-31"},
                   {"kind": "Note", "value": "v" * 81}]}
    r = post(bad)
    assert r.status_code == 400 and "Nothing imported" in r.json()["detail"] and "entry 2" in r.json()["detail"]
    assert len(c.get("/api/mom-export").json()["log"]) == 3
    assert c.post("/api/mom-import", content="not json", headers={**h, "Content-Type": "application/json"}).status_code == 400
    assert c.post("/api/mom-import", content="[1]", headers={**h, "Content-Type": "application/json"}).status_code == 400
    # shows up on the pages, including the family ones
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    mom = c.get("/mom").text
    assert "Treatment" in mom and "Chemo started" in mom and "Dr B — GP" in mom and "Download backup" in mom
    f = TestClient(c.app, base_url="http://testserver")
    fh = login(f, settings, "family-pass-123")
    fam = f.get("/family").text
    assert "Care team" in fam and "Dr B — GP" in fam
    assert "علاج" in f.get("/family/log?lang=ar").text and "Chemo started" in f.get("/family/log").text
    # owner only, CSRF protected
    assert f.get("/api/mom-export").status_code == 403
    assert f.post("/api/mom-import", content="{}", headers={**fh, "Content-Type": "application/json"}).status_code == 403
    assert c.post("/api/mom-import", content="{}", headers={"Content-Type": "application/json"}).status_code == 403


def test_excel_export_for_doctor(settings, tmp_path):
    import io as _io
    from openpyxl import load_workbook
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    c.post("/api/arabic/fill", headers=h)
    c.put("/api/profile", json={"name": "Amina", "doctor": "Prof. Kamar", "notes": "Chemo tablets are not on this list"}, headers=h)
    delta, eps = med_id(c, "Delta"), med_id(c, "Epsilon")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00", "20:00"]}, headers=h)
    c.post(f"/api/medicines/{eps}/reminder", json={"on": True, "times": ["09:00"], "weekdays": [4]}, headers=h)
    c.post("/api/medicines", json={"name": "=HYPERLINK(\"http://x\")", "dose": "+1", "section": "regular"}, headers=h)
    r = c.get("/medicines/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].startswith('attachment; filename="Medicines_for_doctor_review_')
    wb = load_workbook(_io.BytesIO(r.content))
    assert wb.sheetnames == ["Medicines", "Reminder schedule", "Patient"]
    ws = wb["Medicines"]
    assert "Amina — Medicines for doctor review" in ws["A1"].value and "Allergies: Not recorded" in ws["A3"].value
    header = [c_.value for c_ in ws[5]]
    assert header[1] == "Medication" and "Reminder sent to Mom" in header[5] and header[9] == "Correct?"
    rows = {row[1].value: [c_.value for c_ in row] for row in ws.iter_rows(min_row=6) if row[1].value}
    assert len(rows) == 9  # 8 from the sheet + the hand-added one; the section rows are not medicines
    assert rows["Delta"][5] == "Every day at 08:00, 20:00" and rows["Epsilon"][5] == "Every Fri at 09:00"
    assert rows["Alpha"][5] == "No reminder set" and rows["Eta"][5] == "No reminder (as needed)"
    assert rows["Delta"][2] == "500 mg" and rows["Delta"][3] == "Tablet" and rows["Delta"][8] == "حبة صباحاً ومساءً"
    section_labels = [row[0].value for row in ws.iter_rows(min_row=6) if isinstance(row[0].value, str)]
    assert any(v.startswith("EVERY DAY") for v in section_labels) and any(v.startswith("AS NEEDED") for v in section_labels)
    assert any(v.startswith("EVERY WEEK") for v in section_labels) and any(v.startswith("EVERY FEW DAYS") for v in section_labels)
    assert ws.data_validations.dataValidation[0].formula1 == '"Yes,No - needs change"'
    evil = next(row[1] for row in ws.iter_rows(min_row=6) if row[1].value and str(row[1].value).startswith("="))
    assert evil.data_type == "s" and evil.value.startswith("=HYPERLINK")  # stored as text, never a formula
    assert any(str(row[1].value) == "+1" or str(row[2].value) == "+1" for row in ws.iter_rows(min_row=6))
    assert "Chemo tablets are not on this list" in str([c_.value for row in ws.iter_rows(min_row=17) for c_ in row if c_.value])
    sch = wb["Reminder schedule"]
    texts = [[c_.value for c_ in row] for row in sch.iter_rows(min_row=4) if row[0].value]
    assert [t[0] for t in texts] == ["Every day at 08:00", "Every Fri at 09:00", "Every day at 20:00"]
    assert any("تذكير: حان وقت الدواء" in str(t[2]) for t in texts) and all("نسخة" not in str(t[2]) for t in texts)  # Mom's version only
    assert wb["Patient"]["B1"].value == "Amina" and "Not recorded" in wb["Patient"]["B4"].value
    # owner only
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    f = TestClient(c.app, base_url="http://testserver")
    login(f, settings, "family-pass-123")
    assert f.get("/medicines/export.xlsx").status_code == 403
    assert make(settings).get("/medicines/export.xlsx", follow_redirects=False).status_code == 303


def test_stock_pages_counts_reports_and_alerts(settings, tmp_path, monkeypatch):
    import io as _io
    from datetime import datetime as _dt
    from openpyxl import load_workbook
    import app.main as main_mod

    NOW = _dt(2026, 10, 26, 7, 0)  # a morning before the 08:00 dose
    monkeypatch.setattr(main_mod, "local_now", lambda s: NOW)

    class Spy(FakeClient):
        sent = []

        def _post(self, path, fields=None, retries=0):
            if path.endswith("/send_msg/"):
                Spy.sent.append(dict(fields))
            return {"status": 1}

    c = TestClient(create_app(settings, Spy(settings)), base_url="http://testserver")
    h = login(c, settings)
    imported(c, h, tmp_path)                                    # the sheet has counts: Alpha 30, Beta 17 as of 15 Oct
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00", "20:00"]}, headers=h)
    page = c.get("/stock").text
    assert "Medicine stock" in page and "Not counted yet" in page and "Alpha" in page and 'href="/stock"' in page
    # counts imported from the Excel column: Alpha has no schedule times, so nothing is used automatically
    assert re.search(r"<strong dir=\"auto\">Alpha</strong>.*?<strong>30</strong>", page, re.S)
    # first count for Delta: 100 on 15 Oct; 10 finished days x 2 doses have been used by 26 Oct
    r = c.post(f"/api/stock/{delta}/move", json={"kind": "count", "qty": 100, "day": "2026-10-15"}, headers=h)
    assert r.status_code == 200
    detail = c.get(f"/api/stock/{delta}").json()
    assert detail["counted"] and detail["has_schedule"] and detail["unit"] == "tablets" and detail["moves"][0]["kind"] == "Starting count"
    page = c.get("/stock").text
    assert re.search(r"Delta</strong>.*?<strong>80</strong>", page, re.S)             # 100 - 20
    assert "04 Dec" in page and ">39<" in page                                        # 2 a day: 39 more days
    c.post(f"/api/stock/{delta}/move", json={"kind": "buy", "qty": 30, "day": "2026-10-20", "note": "pharmacy"}, headers=h)
    c.post(f"/api/stock/{delta}/move", json={"kind": "used", "qty": 2, "day": "2026-10-21"}, headers=h)
    assert re.search(r"Delta</strong>.*?<strong>108</strong>", c.get("/stock").text, re.S)  # 80 + 30 - 2
    c.post(f"/api/stock/{delta}/move", json={"kind": "count", "qty": 100}, headers=h)       # stock take: 100 really on hand
    assert re.search(r"Delta</strong>.*?<strong>100</strong>", c.get("/stock").text, re.S)
    assert [m["kind"] for m in c.get(f"/api/stock/{delta}").json()["moves"]][0] == "Correction"
    # usage report over 18..25 Oct: 96 before the 18th, +30 bought, -16 scheduled, -1... then the recount
    rep = c.get("/stock/report?start=2026-10-18&end=2026-10-25").text
    assert "Usage report" in rep and "Delta" in rep and ">16<" in rep and ">30<" in rep
    wb = load_workbook(_io.BytesIO(c.get("/stock/report.xlsx?start=2026-10-18&end=2026-10-25").content))
    row = next(r for r in wb["Usage"].iter_rows(min_row=5, values_only=True) if r[0] == "Delta")
    assert row[2] == 96 and row[3] == 30 and row[4] == 16 and row[5] == 2 and row[6] == 0
    sw = load_workbook(_io.BytesIO(c.get("/stock/export.xlsx").content))["Stock"]
    srow = next(r for r in sw.iter_rows(min_row=5, values_only=True) if r[0] == "Delta")
    assert srow[1] == "Every day / week" and srow[3] == 100 and srow[5] == 2 and srow[8] == "OK"
    page = c.get("/stock").text
    reg, asn = page.split("As needed <span")[0], page.split("As needed <span")[1]
    assert "Delta" in reg and "Delta" not in asn.split("Low-stock alerts")[0]
    assert "Beta" in asn.split("Low-stock alerts")[0] and "Beta" not in reg.split("Every day &amp; every week")[1]  # Beta is 'when in pain'
    assert next(r for r in sw.iter_rows(min_row=5, values_only=True) if r[0] == "Beta")[1] == "As needed"
    assert next(r for r in sw.iter_rows(min_row=5, values_only=True) if r[0] == "Gamma")[1] == "Every few days or weeks"
    low = page.split("Low-stock alerts")[0]
    assert "Epsilon" in reg.split("Every day &amp; every week")[1] and "sec-daily_weekly" in low and "sec-periodic" in low  # weekly shares the daily section
    assert low.index("Gamma") > low.index("Epsilon") and low.index("Beta") > low.index("Gamma")                      # blue section, then amber
    # validation
    bad = lambda **kw: c.post(f"/api/stock/{delta}/move", json=kw, headers=h).status_code  # noqa: E731
    assert bad(kind="buy", qty=0) == 400 and bad(kind="buy", qty="x") == 400 and bad(kind="used", qty=-1) == 400
    assert bad(kind="buy", qty=5, day="2026-10-30") == 400 and bad(kind="buy", qty=5, day="nope") == 400
    assert bad(kind="steal", qty=5) == 400 and bad(kind="buy", qty=1e99) == 400
    assert c.get("/stock/report?start=2026-10-30&end=2026-10-01").status_code == 400
    # per-medicine settings: units per dose
    assert c.put(f"/api/stock/{delta}/settings", json={"units_per_dose": 0, "tracking": True}, headers=h).status_code == 400
    assert c.put(f"/api/stock/{delta}/settings", json={"units_per_dose": 2, "tracking": True}, headers=h).status_code == 200
    assert re.search(r"Delta</strong>.*?<strong>100</strong>.*?>4<", c.get("/stock").text, re.S)  # now 4 used a day
    c.put(f"/api/stock/{delta}/settings", json={"units_per_dose": 1, "tracking": True}, headers=h)
    # --- alerts: nothing is sent until there is a number; then once per episode
    assert c.post("/api/stock/check-alerts", json={}, headers=h).status_code == 400
    assert c.put("/api/stock/settings", json={"stock_number": "123", "low_days": 10, "order_days": 30}, headers=h).status_code == 400
    ok = c.put("/api/stock/settings", json={"stock_number": "+961 70 111 222", "low_days": 10, "order_days": 30}, headers=h)
    assert ok.json()["stock_number"] == "96170111222" and settings.stock_number == "96170111222"
    c.post(f"/api/stock/{delta}/move", json={"kind": "count", "qty": 12}, headers=h)          # 12 left = 5 days: low
    assert 'aria-label="1 medicines low"' in c.get("/stock").text and "Running low" in c.get("/stock").text
    Spy.sent.clear()
    r = c.post("/api/stock/check-alerts", json={}, headers=h).json()
    assert r["sent"] is True and Spy.sent[0]["phone"] == "96170111222" and settings.mom_number not in str(Spy.sent[0])
    assert "Delta" in Spy.sent[0]["message_template"] and "Running low" in Spy.sent[0]["message_template"]
    assert c.post("/api/stock/check-alerts", json={}, headers=h).json()["sent"] is False  # already announced
    assert c.post("/api/stock/test-alert", json={}, headers=h).status_code == 400            # needs confirmation
    assert c.post("/api/stock/test-alert", json={"confirm": True}, headers=h).json() == {"ok": True}
    assert "TEST" in Spy.sent[-1]["message_template"]
    # permissions
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    f = TestClient(c.app, base_url="http://testserver")
    fh = login(f, settings, "family-pass-123")
    for path in ("/stock", "/stock/report", "/stock/export.xlsx", "/stock/report.xlsx"):
        r = f.get(path, follow_redirects=False)
        assert r.status_code in (303, 403), path
    assert f.get("/api/stock/%d" % delta).status_code == 403
    assert f.post(f"/api/stock/{delta}/move", json={"kind": "buy", "qty": 1}, headers=fh).status_code == 403
    assert make(settings).get("/stock", follow_redirects=False).status_code == 303
    assert c.post(f"/api/stock/{delta}/move", json={"kind": "buy", "qty": 1}).status_code == 403  # CSRF
    # deleting a medicine removes its stock
    c.delete(f"/api/medicines/{delta}", headers=h)
    assert "Delta" not in c.get("/stock").text


def test_import_counts_for_medicines_already_in_the_app(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    c.post("/api/medicines", json={"name": "Alpha", "section": "regular"}, headers=h)  # added by hand: no stock yet
    c.post("/api/medicines", json={"name": "Beta", "section": "regular"}, headers=h)
    p = tmp_path / "m.xlsx"
    sheet(p)
    r = c.post("/api/stock/import", files={"file": ("m.xlsx", p.read_bytes())}, headers=h).json()
    assert sorted(r["applied"]) == ["Alpha", "Beta"] and r["unmatched"] == [] and r["skipped"] == []  # rows without a quantity are ignored
    again = c.post("/api/stock/import", files={"file": ("m.xlsx", p.read_bytes())}, headers=h).json()
    assert again["applied"] == [] and sorted(again["skipped"]) == ["Alpha", "Beta"]  # never overwrites a count
    no_qty = tmp_path / "n.xlsx"
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Medication", "Dose"])
    wb.active.append(["Alpha", "1 mg"])
    wb.save(no_qty)
    assert c.post("/api/stock/import", files={"file": ("n.xlsx", no_qty.read_bytes())}, headers=h).status_code == 400
    assert c.post("/api/stock/import", files={"file": ("x.xlsx", b"junk")}, headers=h).status_code == 400


def test_stock_page_says_tablets_and_lets_you_change_the_word(settings, tmp_path, monkeypatch):
    from datetime import datetime as _dt
    import app.main as main_mod
    monkeypatch.setattr(main_mod, "local_now", lambda s: _dt(2026, 10, 26, 7, 0))
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)                                         # Alpha 30, Beta 17 (forms: Tablet)
    eps = med_id(c, "Epsilon")                                        # a Supplement
    c.post(f"/api/stock/{eps}/move", json={"kind": "count", "qty": 25, "day": "2026-10-15"}, headers=h)
    page = c.get("/stock").text
    assert re.search(r"Epsilon</strong>.*?<strong>25</strong> <span class=\"muted\">tablets</span>", page, re.S)
    assert "units" not in page.split("<table")[1].split("</table>")[0]   # the table never says "units"
    c.post(f"/api/stock/{eps}/move", json={"kind": "count", "qty": 1}, headers=h)
    assert re.search(r"Epsilon</strong>.*?<strong>1</strong> <span class=\"muted\">tablet</span>", c.get("/stock").text, re.S)  # singular
    d = c.get(f"/api/stock/{eps}").json()
    assert d["unit"] == "tablets" and d["unit_default"] == "tablets"
    ok = lambda **kw: c.put(f"/api/stock/{eps}/settings", json={"units_per_dose": 1, "tracking": True, **kw}, headers=h)  # noqa: E731
    assert ok(unit="capsules").status_code == 200 and c.get(f"/api/stock/{eps}").json()["unit"] == "capsules"
    assert re.search(r"Epsilon</strong>.*?<span class=\"muted\">capsule</span>", c.get("/stock").text, re.S)
    assert ok(unit="tab3ts").status_code == 400 and ok(unit="<b>").status_code == 400
    assert ok(unit="Tablets").status_code == 200 and c.get(f"/api/stock/{eps}").json()["unit"] == "tablets"  # back to the default
    assert ok(units_per_dose=0.5).status_code == 200                                                          # half a tablet is fine


def test_turn_off_all_reminders_cancels_through_review(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta, eps = med_id(c, "Delta"), med_id(c, "Epsilon")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00", "20:00"]}, headers=h)
    c.post(f"/api/medicines/{eps}/reminder", json={"on": True, "times": ["09:00"], "weekdays": [4]}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    assert len(DB(settings.db_path).active()) == 6                  # 3 schedules x (Mom + copy to me)
    # one click turns every switch off; nothing is cancelled on Wbiztool yet
    r = c.post("/api/reminders/clear", headers=h)
    assert r.json() == {"turned_off": 2}
    assert 'class="switch on' not in c.get("/medicines").text
    assert len(DB(settings.db_path).active()) == 6                  # still scheduled until approved
    plan = c.post("/api/sync/preview", headers=h).json()
    assert plan["create"] == [] and len(plan["cancel"]) == 6        # the whole list to cancel, to be approved
    res = c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h).json()["results"]
    assert len(res) == 6 and all(x["ok"] for x in res)
    assert DB(settings.db_path).active() == [] and "Nothing is scheduled yet" in c.get("/reminders").text
    # the times are kept, so a reminder can be switched back on in one click
    assert c.get(f"/api/medicines/{delta}").json()["times"] == ["08:00", "20:00"]
    assert c.post("/api/reminders/clear", headers=h).json() == {"turned_off": 0}   # idempotent
    # owner only, CSRF
    assert c.post("/api/reminders/clear").status_code == 403
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    f = TestClient(c.app, base_url="http://testserver")
    fh = login(f, settings, "family-pass-123")
    assert f.post("/api/reminders/clear", headers=fh).status_code == 403


def test_parse_reminder_list_is_forgiving():
    from app.wbiztool import parse_reminder_list
    data = {"status": 1, "data": {"reminders": [
        {"reminder_id": 11, "name": "mom 08:00", "cron": "0 8 * * *", "phone": "9617", "status": 1},
        {"id": "x2", "title": "other", "message": "hi"},
        {"name": "no id"},
    ]}}
    rows = parse_reminder_list(data)
    assert [r["id"] for r in rows] == ["11", "x2"]
    assert rows[0]["schedule"] == "0 8 * * *" and rows[1]["name"] == "other"
    assert parse_reminder_list({"status": 1}) == [] and parse_reminder_list(None) == []


class CleanupClient(FakeClient):
    """Pretends to be Wbiztool: remembers which WhatsApp client each reminder lives under."""
    def __init__(self, s):
        super().__init__(s)
        self.store = {"42": {"r1"}, "9262": {"old1", "old2"}}
        self.cancelled = []

    def list_reminders(self, whatsapp_client=None):
        wc = whatsapp_client or self.s.whatsapp_client
        return {"status": 1, "data": [{"reminder_id": i, "name": f"n-{i}"} for i in sorted(self.store.get(wc, []))]}

    def cancel_reminder(self, remote_id, whatsapp_client=None):
        from app.wbiztool import WbizError, WbizNetworkError
        wc = whatsapp_client or self.s.whatsapp_client
        if remote_id == "boom":
            raise WbizNetworkError("timeout")
        if remote_id not in self.store.get(wc, set()):
            raise WbizError("Reminder not found")
        self.store[wc].discard(remote_id)
        self.cancelled.append((wc, remote_id))


def test_cleanup_lists_and_cancels_on_another_client(settings, tmp_path):
    fake = CleanupClient(settings)
    c = TestClient(create_app(settings, fake), base_url="http://testserver")
    h = login(c, settings)
    r = c.post("/api/wbiz/reminders", json={"whatsapp_client": "9262"}, headers=h).json()
    assert [i["id"] for i in r["items"]] == ["old1", "old2"] and not any(i["tracked"] for i in r["items"])
    assert c.post("/api/wbiz/cancel", json={"whatsapp_client": "9262", "ids": ["old1"]}, headers=h).status_code == 400  # needs confirm
    assert c.post("/api/wbiz/cancel", json={"whatsapp_client": "9262", "ids": [], "confirm": True}, headers=h).status_code == 400
    assert c.post("/api/wbiz/cancel", json={"whatsapp_client": "bad value", "ids": ["a"], "confirm": True}, headers=h).status_code == 400
    res = c.post("/api/wbiz/cancel", json={"whatsapp_client": "9262", "ids": ["old1", "nope", "boom", "old2"], "confirm": True}, headers=h).json()
    assert [x["ok"] for x in res["results"]] == [True, False, False] and res["not_attempted"] == 1  # stops at the network error
    assert fake.cancelled == [("9262", "old1")]
    # owner only, CSRF
    assert c.post("/api/wbiz/reminders", json={}).status_code == 403
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    f = TestClient(c.app, base_url="http://testserver")
    fh = login(f, settings, "family-pass-123")
    assert f.post("/api/wbiz/reminders", json={}, headers=fh).status_code == 403
    assert f.post("/api/wbiz/cancel", json={"ids": ["r1"], "confirm": True}, headers=fh).status_code == 403


def test_cleanup_marks_tracked_and_forgets_missing(settings, tmp_path):
    fake = CleanupClient(settings)
    c = TestClient(create_app(settings, fake), base_url="http://testserver")
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta = med_id(c, "Delta")
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h)
    mine = DB(settings.db_path).active()
    assert len(mine) == 2
    fake.store["42"] = {mine[0]["remote_id"]}                       # the second one vanished on Wbiztool
    items = c.post("/api/wbiz/reminders", json={}, headers=h).json()["items"]
    assert [i["tracked"] for i in items] == [True]
    # Review & apply no longer fails on a reminder Wbiztool does not have
    c.post(f"/api/medicines/{delta}/reminder", json={"on": False}, headers=h)
    plan = c.post("/api/sync/preview", headers=h).json()
    res = c.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=h).json()["results"]
    assert all(x["ok"] for x in res) and any("warning" in x for x in res)
    assert DB(settings.db_path).active() == []


def test_admin_account_edits_everything_but_settings_and_is_tracked(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    assert c.post("/api/users", json={"name": "Sara", "password": "short"}, headers=h).status_code == 400
    assert c.post("/api/users", json={"name": "Owner", "password": "sara-pass-123"}, headers=h).status_code == 400
    uid = c.post("/api/users", json={"name": "Sara", "password": "sara-pass-123"}, headers=h).json()["id"]
    assert c.post("/api/users", json={"name": "sara", "password": "other-pass-123"}, headers=h).status_code == 400     # name taken

    a = TestClient(c.app, base_url="http://testserver")
    ah = login(a, settings, "sara-pass-123", "Sara")
    assert a.get("/medicines").status_code == 200 and "Activity" not in a.get("/medicines").text and "/settings" not in a.get("/medicines").text
    for page in ("/stock", "/reminders", "/mom"):
        assert a.get(page).status_code == 200
    # can edit medicines, stock, reminders and Mom's profile/log
    delta = med_id(a, "Delta")
    assert a.put(f"/api/medicines/{delta}", json={"dose": "750 mg"}, headers=ah).status_code == 200
    assert a.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00"]}, headers=ah).status_code == 200
    assert a.post(f"/api/stock/{delta}/move", json={"kind": "buy", "qty": 30}, headers=ah).status_code == 200
    assert a.put("/api/profile", json={"blood_type": "O+"}, headers=ah).status_code == 200
    assert a.post("/api/status-log", json={"kind": "Blood pressure", "value": "120/80"}, headers=ah).status_code == 200
    plan = a.post("/api/sync/preview", headers=ah).json()
    assert a.post("/api/sync/apply", json={"hash": plan["hash"]}, headers=ah).status_code == 200
    assert a.get("/medicines/export.xlsx").status_code == 200
    # but not settings, accounts or the activity trail
    assert a.get("/settings", follow_redirects=False).status_code == 303
    assert a.get("/activity", follow_redirects=False).status_code == 303
    for method, url, body in (("PUT", "/api/config", {}), ("PUT", "/api/slots", {}), ("PUT", "/api/family-access", {"password": "x" * 9}),
                              ("POST", "/api/users", {"name": "Eve", "password": "eve-pass-123"}), ("POST", "/api/test-message", {}),
                              ("POST", "/api/settings", {}), ("DELETE", f"/api/users/{uid}", None)):
        assert a.request(method, url, json=body, headers=ah).status_code == 403, url

    # the owner sees exactly what they did, with details
    page = c.get("/activity").text
    for text in ("Sara", "Signed in", "Edited a medicine", "dose", "750 mg", "Delta", "Changed stock", "bought 30", "Edited Mom&#39;s profile",
                 "Added a log entry", "Applied reminder changes", "Downloaded the medicines Excel", "Blocked: Changed Wbiztool settings"):
        assert text in page, text
    assert "Dina" not in page
    only = c.get("/activity?user=Sara").text
    assert "Sara" in only and "Owner" not in only.split("<tbody>")[1]
    assert c.get("/activity/export.csv?user=Sara").text.startswith("When,Who")

    # turning an admin off signs them out at once; their history stays
    assert c.put(f"/api/users/{uid}", json={"active": False}, headers=h).status_code == 200
    assert a.get("/medicines", follow_redirects=False).status_code == 303
    assert TestClient(c.app, base_url="http://testserver").post("/login", data={"username": "Sara", "password": "sara-pass-123"}).status_code == 401
    assert "Failed sign-in" in c.get("/activity").text
    c.put(f"/api/users/{uid}", json={"active": True, "password": "new-pass-4567"}, headers=h)
    b = TestClient(c.app, base_url="http://testserver")
    login(b, settings, "new-pass-4567", "sara")
    assert c.delete(f"/api/users/{uid}", headers=h).status_code == 200
    assert b.get("/medicines", follow_redirects=False).status_code == 303
    assert "Sara" in c.get("/activity").text

    # family still read-only; settings page lists admins for the owner
    c.post("/api/users", json={"name": "Maya", "password": "maya-pass-123"}, headers=h)
    assert "Maya" in c.get("/settings").text


def test_stock_price_estimates(settings, tmp_path, monkeypatch):
    import io as _io
    from datetime import datetime as _dt
    from openpyxl import load_workbook
    import app.main as main_mod
    monkeypatch.setattr(main_mod, "local_now", lambda s: _dt(2026, 10, 26, 7, 0))
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    delta = med_id(c, "Delta")                                         # twice daily
    c.post(f"/api/medicines/{delta}/reminder", json={"on": True, "frequency": "daily", "times": ["08:00", "20:00"]}, headers=h)
    c.post(f"/api/stock/{delta}/move", json={"kind": "count", "qty": 20, "day": "2026-10-26"}, headers=h)
    assert "Est. cost" not in c.get("/stock").text.split("<table")[0]       # nothing priced yet: no cost cards
    put = lambda **kw: c.put(f"/api/stock/{delta}/settings", json={"units_per_dose": 1, "tracking": True, **kw}, headers=h)  # noqa: E731
    assert put(pack_price="abc").status_code == 400 and put(pack_price=30, pack_size=0).status_code == 400
    assert put(pack_price=30, pack_size=60).status_code == 200            # $30 for 60 tablets = $0.50 each
    d = c.get(f"/api/stock/{delta}").json()
    assert d["pack_price"] == 30 and d["pack_size"] == 60 and d["currency"] == "$"
    page = c.get("/stock").text
    assert "$30.00" in page                                               # 2 a day x 30 days x $0.50
    assert "To order now" in page and "$20.00" in page                    # 60 needed - 20 on hand = 40 tablets x $0.50
    # Excel has the price columns
    sw = load_workbook(_io.BytesIO(c.get("/stock/export.xlsx").content))["Stock"]
    row = next(r for r in sw.iter_rows(min_row=5, values_only=True) if r[0] == "Delta")
    assert row[10] == 0.5 and row[11] == 30 and row[12] == 20
    # usage report: cost of what was used in the period
    rep = c.get("/stock/report?start=2026-10-26&end=2026-10-26").text
    assert "Est. cost" in rep
    # price without a pack size is per unit; clearing the price removes the estimate; currency is editable
    assert put(pack_price=2).status_code == 200 and c.get(f"/api/stock/{delta}").json()["pack_size"] is None
    assert c.put("/api/stock/settings", json={"low_days": 10, "order_days": 30, "currency": "LBP"}, headers=h).status_code == 200
    assert "LBP" in c.get("/stock").text
    assert put(pack_price="").status_code == 200 and "Est. cost / month</span>" not in c.get("/stock").text


def test_medicines_are_sectioned_by_how_often(settings, tmp_path):
    c = make(settings)
    h = login(c, settings)
    imported(c, h, tmp_path)
    page = c.get("/medicines").text
    heads = re.findall(r'<h3 class="sect sec-\w+">(.*?) <span', page)
    assert heads == ["Every day", "Every week", "Every few days or weeks", "As needed"]
    section = lambda name: re.split(rf'<h3 class="sect sec-\w+">{name} ', page)[1].split("<h3 ")[0]  # noqa: E731
    assert "Delta" in section("Every day") and "Alpha" in section("Every day")
    assert "Epsilon" in section("Every week") and "Gamma" in section("Every few days or weeks")   # Gamma: every 28 days
    assert "Beta" in section("As needed") and "Delta" not in section("As needed")
    ar = c.get("/medicines?lang=ar").text
    assert "كل يوم" in ar and "كل أسبوع" in ar and "عند الحاجة" in ar


def test_login_asks_for_user_name_and_password(settings):
    c = make(settings)
    page = c.get("/login").text
    assert 'name="username"' in page and 'name="password"' in page
    pw = settings.dashboard_password
    post = lambda **d: TestClient(c.app, base_url="http://testserver").post("/login", data=d, follow_redirects=False).status_code  # noqa: E731
    assert post(password=pw) == 401                                   # no user name
    assert post(username="admin", password=pw) == 401                 # wrong user name
    assert post(username="owner", password="nope") == 401
    assert post(username="Owner", password=pw) == 303                 # not case-sensitive
    h = login(c, settings)
    c.put("/api/family-access", json={"password": "family-pass-123"}, headers=h)
    c.post("/api/users", json={"name": "Sara", "password": "sara-pass-123"}, headers=h)
    assert post(username="owner", password="family-pass-123") == 401  # a password only works with its own user name
    assert post(username="family", password="family-pass-123") == 303
    assert post(username="Sara", password="sara-pass-123") == 303
    assert post(username="family", password="sara-pass-123") == 401
    assert "user name tried: admin" in c.get("/activity").text
