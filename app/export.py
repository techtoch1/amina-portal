"""Excel export of the medicines in the dashboard, laid out for a doctor to review and confirm."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .excel import classify_frequency, load_from_db, norm
from .schedule import build_desired, cron_human, next_occurrence

BRAND = "1D7A6C"
HEAD = PatternFill("solid", fgColor=BRAND)
SECTION = PatternFill("solid", fgColor="E7F3F0")
ASK = PatternFill("solid", fgColor="FFF4DF")
THIN = Side(style="thin", color="C9D6D2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
RTL = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
_DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _put(ws, row, col, value, **style):
    """Write text safely: a value starting with = + - @ must never become a formula."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def reminder_text(m, now):
    """What the app sends for this medicine, in words."""
    if m is None or not m.remind:
        return "No reminder (as needed)" if m is not None and m.kind == "as_needed" else "No reminder set"
    times = ", ".join(m.times)
    if m.kind == "daily":
        return f"Every day at {times}"
    if m.kind == "weekly":
        return f"Every {', '.join(_DAYS[d] for d in m.weekdays)} at {times}"
    if m.kind == "every_n":
        nxt = next_occurrence(m, now)
        return f"Every {m.every_n} days at {times}" + (f" (next: {nxt.strftime('%d %b %Y')})" if nxt else "")
    return "No reminder set"


def build_doctor_workbook(db, settings, now):
    res, _stopped = load_from_db(db)
    parsed = {m.row: m for m in res.medicines}
    profile = db.profile()
    name = profile.get("name") or "Patient"
    wb = Workbook()

    # ---------------- sheet 1: the medicines ----------------
    ws = wb.active
    ws.title = "Medicines"
    cols = ["#", "Medication", "Dose", "Form", "When to take", "Reminder sent to Mom (WhatsApp)", "For", "Notes",
            "Arabic instruction (sent in the reminder)", "Correct?", "Doctor's correction / comment"]
    widths = [5, 26, 12, 14, 30, 30, 24, 36, 32, 14, 40]
    _put(ws, 1, 1, f"{name} — Medicines for doctor review", font=Font(bold=True, size=15))
    _put(ws, 2, 1, f"Prepared {now.strftime('%d %b %Y')} from the reminder dashboard. Please confirm each line "
                    "(dose, timing and reminder times) or write the correction." + (
                        f" Prescriber: {profile['doctor']}." if profile.get("doctor") else ""),
         font=Font(italic=True, color="586F69"))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws.cell(row=2, column=1).alignment = WRAP
    ws.row_dimensions[2].height = 32
    allergies = profile.get("allergies") or "Not recorded — please confirm"
    _put(ws, 3, 1, f"Allergies: {allergies}", font=Font(bold=True, color="B3261E" if profile.get("allergies") else "9A6700"))
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols))

    header_row = 5
    for c, (title, width) in enumerate(zip(cols, widths), 1):
        _put(ws, header_row, c, title, font=Font(bold=True, color="FFFFFF"), fill=HEAD, border=BOX,
             alignment=Alignment(wrap_text=True, vertical="center"))
        ws.column_dimensions[chr(64 + c)].width = width
    ws.row_dimensions[header_row].height = 34

    groups = {"Regular": [], "As needed": []}
    for r in sorted(db.list_medicines(), key=lambda x: x["id"]):
        if not r["active"]:
            continue
        kind, _n, _p = classify_frequency(norm(r["frequency"]))
        groups["As needed" if kind == "as_needed" else "Regular"].append(r)

    row, number = header_row + 1, 0
    validation = DataValidation(type="list", formula1='"Yes,No - needs change"', allow_blank=True)
    ws.add_data_validation(validation)
    for section, items in groups.items():
        if not items:
            continue
        _put(ws, row, 1, f"{section.upper()} ({len(items)})", font=Font(bold=True), fill=SECTION)
        for c in range(2, len(cols) + 1):
            ws.cell(row=row, column=c).fill = SECTION
        row += 1
        for r in items:
            number += 1
            m = parsed.get(r["id"])
            values = [number, r["name"], r["dose"], r["form"], r["when_en"], reminder_text(m, now), r["purpose"],
                      r["notes_en"] or r["notes_ar"], r["instruction"], None, None]
            for c, value in enumerate(values, 1):
                cell = _put(ws, row, c, value, border=BOX, alignment=RTL if c == 9 else WRAP)
                if c in (10, 11):
                    cell.fill = ASK
            validation.add(ws.cell(row=row, column=10))
            row += 1
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    if profile.get("notes"):
        row += 1
        _put(ws, row, 1, "Also: " + profile["notes"], font=Font(italic=True))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(cols))
        ws.cell(row=row, column=1).alignment = WRAP
        ws.row_dimensions[row].height = 30
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"

    # ---------------- sheet 2: what Mom actually receives ----------------
    ws2 = wb.create_sheet("Reminder schedule")
    _put(ws2, 1, 1, "What Mom receives on WhatsApp (medicines due at the same time are grouped in one message)",
         font=Font(bold=True, size=13))
    for c, (title, width) in enumerate(zip(["Schedule", "Medicines in the message", "Message (Arabic)", "Correct?", "Comment"],
                                           [26, 38, 60, 14, 36]), 1):
        _put(ws2, 3, c, title, font=Font(bold=True, color="FFFFFF"), fill=HEAD, border=BOX)
        ws2.column_dimensions[chr(64 + c)].width = width
    desired = build_desired([m for m in res.medicines if m.remind], now, settings.timezone, [("mom", "x")],
                            with_purpose=db.get_setting("msg.purpose", "1") == "1")
    r2 = 4
    for d in desired:
        when = cron_human(d.cron) if d.kind == "cron" else f"Once on {d.fire_at[8:10]}/{d.fire_at[5:7]}/{d.fire_at[:4]} at {d.fire_at[11:]}"
        for c, value in enumerate([when, ", ".join(d.names), d.message, None, None], 1):
            cell = _put(ws2, r2, c, value, border=BOX, alignment=RTL if c == 3 else WRAP)
            if c >= 4:
                cell.fill = ASK
        r2 += 1
    if not desired:
        _put(ws2, 4, 1, "No reminders are switched on yet.", font=Font(italic=True))
    dv2 = DataValidation(type="list", formula1='"Yes,No - needs change"', allow_blank=True)
    ws2.add_data_validation(dv2)
    if desired:
        dv2.add(f"D4:D{3 + len(desired)}")
    ws2.freeze_panes = "A4"
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- sheet 3: patient details to confirm ----------------
    ws3 = wb.create_sheet("Patient")
    rows = [("Name", profile.get("name")), ("Age / date of birth", profile.get("age")), ("Conditions", profile.get("conditions")),
            ("Allergies", profile.get("allergies") or "Not recorded — please confirm"), ("Doctor", profile.get("doctor")),
            ("Doctor / clinic phone", profile.get("doctor_phone")), ("Pharmacy", profile.get("pharmacy")),
            ("Emergency contact", profile.get("emergency_contact"))]
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 90
    for i, (label, value) in enumerate(rows, 1):
        _put(ws3, i, 1, label, font=Font(bold=True), border=BOX, alignment=WRAP)
        _put(ws3, i, 2, value or "—", border=BOX, alignment=WRAP)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------- stock reports ----------------
def _sheet(wb, title, heading, subtitle, columns, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    _put(ws, 1, 1, heading, font=Font(bold=True, size=15))
    _put(ws, 2, 1, subtitle, font=Font(italic=True, color="586F69"))
    for c, (name, width) in enumerate(zip(columns, widths), 1):
        _put(ws, 4, c, name, font=Font(bold=True, color="FFFFFF"), fill=HEAD, border=BOX,
             alignment=Alignment(wrap_text=True, vertical="center"))
        ws.column_dimensions[chr(64 + c)].width = width
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


_STATUS = {"ok": "OK", "low": "LOW", "out": "OUT", "none": "Not counted"}
_FILL = {"low": PatternFill("solid", fgColor="FFF4DF"), "out": PatternFill("solid", fgColor="FDECEA")}


def build_stock_workbook(rows, low_days, order_days, now, profile):
    from .stock import num

    wb = Workbook()
    name = profile.get("name") or "Patient"
    ws = _sheet(wb, "Stock", f"{name} — Medicine stock",
                f"As of {now.strftime('%d %b %Y %H:%M')}. Estimated from the schedule plus manual changes; "
                f"LOW = under {low_days} days left.",
                ["Medication", "Form", "On hand", "Unit", "Use / day", "Days left", "Runs out", "Status", f"Order for {order_days} days"],
                [30, 14, 11, 12, 11, 11, 14, 14, 18], first=True)
    for i, r in enumerate(rows, 5):
        values = [r["name"], r["form"], r["on_hand"] if r["tracked"] else None, r["unit"] if r["tracked"] else "",
                  round(r["daily"], 2) if r["tracked"] and r["daily"] else None,
                  r["days_left"] if r["tracked"] else None,
                  r["run_out"].strftime("%d %b %Y") if r["run_out"] else "", _STATUS[r["status"]],
                  r["order_qty"] or None]
        for c, v in enumerate(values, 1):
            cell = _put(ws, i, c, v, border=BOX, alignment=WRAP)
            if r["status"] in _FILL:
                cell.fill = _FILL[r["status"]]
    n = 5 + len(rows) + 1
    todo = [r for r in rows if r["order_qty"] and r["status"] in ("low", "out")]
    _put(ws, n, 1, "To order" + (": " + ", ".join(f"{r['name']} × {r['order_qty']}" for r in todo) if todo else ": nothing right now"),
         font=Font(bold=True))
    return _bytes(wb)


def build_usage_workbook(rows, start, end, profile):
    wb = Workbook()
    name = profile.get("name") or "Patient"
    days = (end - start).days + 1
    ws = _sheet(wb, "Usage", f"{name} — Medicine usage",
                f"{start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')} ({days} days). Scheduled use is an estimate from the reminder schedule.",
                ["Medication", "Unit", "Opening", "Bought / added", "Scheduled use", "Other use", "Corrections", "Closing", "Avg used / day"],
                [30, 12, 11, 15, 15, 12, 13, 11, 15], first=True)
    for i, r in enumerate(rows, 5):
        for c, v in enumerate([r["name"], r["unit"], r["opening"], r["bought"], r["auto"], r["manual"], r["adjust"], r["closing"], r["per_day"]], 1):
            _put(ws, i, c, v, border=BOX, alignment=WRAP)
    n = 5 + len(rows)
    _put(ws, n, 1, "Total used", font=Font(bold=True))
    _put(ws, n, 5, round(sum(r["auto"] for r in rows), 2), font=Font(bold=True))
    _put(ws, n, 6, round(sum(r["manual"] for r in rows), 2), font=Font(bold=True))
    return _bytes(wb)


def _bytes(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
